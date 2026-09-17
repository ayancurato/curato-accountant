import pytest
import zipfile
import io
import openpyxl
import os
from datetime import date
from decimal import Decimal
from app.models import Transaction, Document
from uuid import uuid4

from tests.conftest import TestingSessionLocal

@pytest.fixture
def db_session():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()

# ---------------------------------------------------------
# Helper to create mock transactions with anomalies and docs
# ---------------------------------------------------------
def create_mock_tx(db, type="EXPENSE", status="APPROVED", anomalies=None, doc=True, 
                   date_val=None, amount=1000, tds=False, is_asset=False, gst=180):
    from app.models import Company, User
    company = db.query(Company).first()
    user = db.query(User).first()
    
    doc_id = None
    if doc:
        d = Document(
            company_id=company.id,
            uploaded_by=user.id,
            file_name="test_doc.pdf",
            file_path="tests/mock_file.pdf",
            file_type="application/pdf",
            file_size=1024,
            status="PROCESSED",
            file_hash=str(uuid4())
        )
        db.add(d)
        db.commit()
        doc_id = d.id

    tx = Transaction(
        company_id=company.id,
        type=type,
        status=status,
        date=date_val or date(2030, 1, 15),
        total_amount=Decimal(amount),
        net_amount=Decimal(amount) - Decimal(gst),
        gst_amount=Decimal(gst),
        vendor_customer="Test Entity",
        anomalies=anomalies,
        document_id=doc_id,
        expense_type="ASSET" if is_asset else "OPERATING_EXPENSE",
        tds_applicable=tds,
        tds_amount=Decimal(100) if tds else Decimal(0)
    )
    db.add(tx)
    db.commit()
    return tx

# ---------------------------------------------------------
# Readiness Tests
# ---------------------------------------------------------
def test_readiness_zero_transactions(client, auth_headers, db_session):
    resp = client.get("/api/v1/ca/readiness?start_date=2040-01-01&end_date=2040-12-31", headers=auth_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["transaction_count"] == 0
    assert data["needs_review_count"] == 0
    assert data["missing_documents_count"] == 0
    assert data["potential_duplicates_count"] == 0
    assert data["gst_issues_count"] == 0

def test_readiness_counts_and_exceptions(client, auth_headers, db_session):
    # Missing Doc
    create_mock_tx(db_session, date_val=date(2035, 1, 1), doc=False)
    # Needs Review
    create_mock_tx(db_session, date_val=date(2035, 1, 2), status="NEEDS_REVIEW")
    # Duplicate
    create_mock_tx(db_session, date_val=date(2035, 1, 3), anomalies=[{"type": "DUPLICATE_TRANSACTION"}])
    # GST Issue
    create_mock_tx(db_session, date_val=date(2035, 1, 4), anomalies=[{"type": "MISSING_GSTIN"}])
    # GST Issue 2
    create_mock_tx(db_session, date_val=date(2035, 1, 5), anomalies=[{"type": "TOTAL_MISMATCH"}])
    # Normal
    create_mock_tx(db_session, date_val=date(2035, 1, 6))

    resp = client.get("/api/v1/ca/readiness?start_date=2035-01-01&end_date=2035-12-31", headers=auth_headers)
    assert resp.status_code == 200
    data = resp.json()
    
    assert data["transaction_count"] == 6
    assert data["missing_documents_count"] == 1
    assert data["needs_review_count"] == 1
    assert data["potential_duplicates_count"] == 1
    assert data["gst_issues_count"] == 2
    
    # Check Exception Transaction Details format
    assert len(data["gst_issues"]) == 2
    assert "id" in data["gst_issues"][0]
    assert "date" in data["gst_issues"][0]
    assert data["gst_issues"][0]["type"] == "EXPENSE"

# ---------------------------------------------------------
# Export Tests
# ---------------------------------------------------------
def test_export_empty(client, auth_headers):
    resp = client.get("/api/v1/ca/export?start_date=2040-01-01&end_date=2040-12-31", headers=auth_headers)
    assert resp.status_code == 200
    assert "CA_Pack_" in resp.headers["content-disposition"]
    
    zip_buffer = io.BytesIO(resp.content)
    with zipfile.ZipFile(zip_buffer, 'r') as zf:
        assert "CA_Pack.xlsx" in zf.namelist()

def test_export_full_workbook_and_exceptions(client, auth_headers, db_session):
    # Setup mock files
    os.makedirs("tests", exist_ok=True)
    with open("tests/mock_file.pdf", "w") as f:
        f.write("mock content")
        
    # Approved Income
    create_mock_tx(db_session, type="INCOME", status="APPROVED", date_val=date(2037, 1, 1), amount=5000, tds=True)
    # Approved Expense (Asset)
    create_mock_tx(db_session, type="EXPENSE", status="APPROVED", date_val=date(2037, 1, 2), is_asset=True)
    # Unapproved Transaction (Should be excluded from ledgers, but included in exceptions if it has one)
    create_mock_tx(db_session, type="EXPENSE", status="DRAFT", date_val=date(2037, 1, 3), anomalies=[{"type": "UNKNOWN"}])
    # Missing Doc
    create_mock_tx(db_session, type="EXPENSE", status="APPROVED", date_val=date(2037, 1, 4), doc=False)
    # Duplicate
    create_mock_tx(db_session, type="EXPENSE", status="APPROVED", date_val=date(2037, 1, 5), anomalies=[{"type": "DUPLICATE_TRANSACTION"}])

    resp = client.get("/api/v1/ca/export?start_date=2037-01-01&end_date=2037-12-31", headers=auth_headers)
    assert resp.status_code == 200
    
    zip_buffer = io.BytesIO(resp.content)
    with zipfile.ZipFile(zip_buffer, 'r') as zf:
        files = zf.namelist()
        
        # Check ZIP structure
        assert "CA_Pack.xlsx" in files
        assert "exceptions/missing_documents.csv" in files
        assert "exceptions/needs_review.csv" in files
        assert "exceptions/duplicates.csv" in files
        assert "exceptions/gst_issues.csv" in files
        assert "exceptions/other_flagged.csv" in files
        
        # Check Docs
        docs = [f for f in files if f.startswith("docs/")]
        assert len(docs) > 0 # At least 4 documents attached to the txs above
        
        # Check Excel Sheets
        xlsx_data = zf.read("CA_Pack.xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_data))
        
        assert "Income Ledger" in wb.sheetnames
        assert "Expense Ledger" in wb.sheetnames
        assert "Asset Transactions" in wb.sheetnames
        assert "Category-wise Expenses" in wb.sheetnames
        assert "Vendor-wise Expenses" in wb.sheetnames
        assert "Customer-wise Income" in wb.sheetnames
        assert "Monthly Summary" in wb.sheetnames
        assert "GST Summary" in wb.sheetnames
        assert "TDS Summary" in wb.sheetnames
        assert "Complete Ledger" in wb.sheetnames
        
        # Verify Unapproved transaction is EXCLUDED from final ledgers
        ws_complete = wb["Complete Ledger"]
        rows = list(ws_complete.rows)
        # Header + 4 approved transactions
        assert len(rows) == 5 
        statuses = [row[2].value for row in rows[1:]]
        assert "DRAFT" not in statuses
        
        # Verify Asset Transactions
        ws_asset = wb["Asset Transactions"]
        # Header + 1 asset
        assert len(list(ws_asset.rows)) == 2
        
        # Verify TDS Summary
        ws_tds = wb["TDS Summary"]
        # Header + 1 tds applicable
        assert len(list(ws_tds.rows)) == 2
        
        # Check Exceptions
        other = zf.read("exceptions/other_flagged.csv").decode("utf-8")
        assert "DRAFT" in other # Unapproved transaction is still in exception reports
        
        missing = zf.read("exceptions/missing_documents.csv").decode("utf-8")
        assert len(missing.strip().split("\n")) == 2 # Header + 1 missing doc
